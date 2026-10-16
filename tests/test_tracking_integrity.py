import io
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from export.exporter import HEADERS, export_jobs_to_xlsx
from export.workbook_appender import append_jobs_to_workbook
from process_excel_links import APPLICATION_HEADERS, INPUT_HEADERS, append_application, process_workbook


def tracker():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Applications"
    sheet.append(HEADERS + ["Notes"])
    sheet.append([
        date(2026, 10, 1), "Example", "Analyst",
        "https://example.com/jobs/1?utm_source=linkedin", "Interview",
        "New York, NY", "Hybrid", "$80,000/yr", date(2026, 10, 25),
        "LinkedIn", "=1+1",
    ])
    return workbook


def workbook_bytes(workbook):
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def job(**overrides):
    return {
        "company": "Example", "job_title": "Updated analyst",
        "job_link": "https://www.example.com/jobs/1/?utm_source=indeed",
        "location": "New York, NY", "work_type": "Hybrid", "salary": "$85,000/yr",
        **overrides,
    }


def test_excel_duplicate_update_preserves_tracking_and_existing_formulas(tmp_path):
    path, summary = append_jobs_to_workbook(
        workbook_bytes(tracker()), "tracker.xlsx",
        [job(date_applied="10/16/2026", status="", follow_up="", found_on="")],
        outdir=tmp_path, duplicate_mode="update",
    )
    sheet = load_workbook(path).active

    assert summary["updated"] == 1
    assert summary["added"] == 0
    assert sheet.cell(2, 1).value.date() == date(2026, 10, 1)
    assert sheet.cell(2, 5).value == "Interview"
    assert sheet.cell(2, 9).value.date() == date(2026, 10, 25)
    assert sheet.cell(2, 10).value == "LinkedIn"
    assert sheet.cell(2, 8).value == "$85,000/yr"
    assert sheet.cell(2, 11).value == "=1+1"
    assert sheet.cell(2, 11).data_type == "f"


def test_excel_update_allows_explicit_tracking_edits(tmp_path):
    path, _ = append_jobs_to_workbook(
        workbook_bytes(tracker()), "tracker.xlsx",
        [job(status="Offer", follow_up="10/28/2026", found_on="Referral")],
        outdir=tmp_path, duplicate_mode="update",
    )
    sheet = load_workbook(path).active

    assert sheet.cell(2, 5).value == "Offer"
    assert sheet.cell(2, 9).value == "10/28/2026"
    assert sheet.cell(2, 10).value == "Referral"


def test_excel_skip_recognizes_the_same_posting_with_different_tracking_parameters(tmp_path):
    path, summary = append_jobs_to_workbook(
        workbook_bytes(tracker()), "tracker.xlsx", [job()], outdir=tmp_path,
    )

    assert summary["skipped"] == 1
    assert summary["added"] == 0
    assert load_workbook(path).active.max_row == 2


@pytest.mark.parametrize("new_link", (
    "https://example.com/jobs/2?utm_source=linkedin",
    "https://example.com/jobs/1?id=2",
    "https://example.com/Jobs/1",
))
def test_excel_deduplication_preserves_real_job_identity(tmp_path, new_link):
    _, summary = append_jobs_to_workbook(
        workbook_bytes(tracker()), "tracker.xlsx", [job(job_link=new_link)], outdir=tmp_path,
    )

    assert summary["added"] == 1
    assert summary["skipped"] == 0


@pytest.mark.parametrize("workflow", ("export", "append", "legacy_cli"))
def test_job_text_is_saved_as_literal_strings_in_every_excel_workflow(tmp_path, workflow):
    data = job(
        company="=1+1", job_title="=2+2", status="=3+3",
        location="=4+4", salary="=5+5", found_on="=6+6", follow_up="=7+7",
    )
    if workflow == "export":
        sheet = load_workbook(export_jobs_to_xlsx([data], outdir=tmp_path)).active
    elif workflow == "append":
        workbook = Workbook()
        workbook.active.title = "Applications"
        workbook.active.append(HEADERS)
        path, _ = append_jobs_to_workbook(workbook_bytes(workbook), "tracker.xlsx", [data], outdir=tmp_path)
        sheet = load_workbook(path).active
    else:
        workbook = Workbook()
        workbook.active.title = "Applications"
        workbook.active.append(APPLICATION_HEADERS)
        append_application(workbook.active, data, data["job_link"])
        sheet = load_workbook(workbook_bytes(workbook)).active

    for column in (2, 3, 5, 6, 8, 9, 10):
        assert sheet.cell(2, column).value.startswith("=")
        assert sheet.cell(2, column).data_type == "s"


def test_legacy_cli_skips_tracking_variants_without_rescraping(tmp_path, monkeypatch):
    workbook = tracker()
    inputs = workbook.create_sheet("Input")
    inputs.append(INPUT_HEADERS)
    inputs.append([job()["job_link"], "", "Pending", "", ""])
    path = tmp_path / "tracker.xlsx"
    workbook.save(path)

    def unexpected_scrape(*args, **kwargs):
        pytest.fail("A duplicate posting should not be scraped")

    monkeypatch.setattr("scraper.browser_scraper_v2.parse_job_with_browser", unexpected_scrape)

    assert process_workbook(path) == (0, 1, 0)
    output = load_workbook(path)
    assert output["Applications"].max_row == 2
    assert output["Input"].cell(2, 3).value == "Duplicate"
